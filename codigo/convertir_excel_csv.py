from pathlib import Path
BASE = Path(__file__).resolve().parent.parent
import csv
from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook

carpeta = BASE / "datos"
archivo_excel = carpeta / "Online Retail.xlsx"
archivo_csv = carpeta / "online_retail.csv"

print("Iniciando conversion de Excel a CSV...")

wb = load_workbook(
    archivo_excel,
    read_only=True,
    data_only=True
)

hoja = wb.active
total = 0

with open(archivo_csv, "w", newline="", encoding="utf-8") as f:
    escritor = csv.writer(f)

    for fila in hoja.iter_rows(values_only=True):
        valores = [
            valor.isoformat(sep=" ")
            if isinstance(valor, datetime)
            else valor.isoformat()
            if isinstance(valor, date)
            else valor
            for valor in fila
        ]

        escritor.writerow(valores)
        total += 1

        if total % 100000 == 0:
            print(f"Filas procesadas: {total}")

wb.close()

print("\nConversion finalizada.")
print(f"Filas totales incluyendo encabezado: {total}")
print(f"Registros de datos: {total - 1}")
print(f"Archivo generado: {archivo_csv}")
